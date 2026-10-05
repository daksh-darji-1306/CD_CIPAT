import os
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Please install python-docx: pip install python-docx")
    exit(1)

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches, Pt as PptxPt
except ImportError:
    print("Please install python-pptx: pip install python-pptx")
    exit(1)

def generate_report():
    doc = Document()
    
    # Title Page
    title = doc.add_heading('Mini C Compiler', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Compiler Design (3170701) - Project Report').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('C. K. Pithawala College of Engineering and Technology').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # Section 1: Objective
    doc.add_heading('1. Objective', level=1)
    doc.add_paragraph("The objective of this project is to implement a complete, end-to-end Mini C Compiler that translates a subset of the C programming language into executable output, demonstrating all phases of the compilation process including lexical analysis, syntax analysis, semantic analysis, intermediate code generation, optimization, and execution via a virtual machine.")

    # Section 2: Architecture
    doc.add_heading('2. Architecture', level=1)
    doc.add_paragraph("The architecture is split into a frontend and a backend:")
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Backend: ").bold = True
    ul.add_run("A Python-based FastAPI service that runs the compiler phases entirely from scratch without external parsing libraries.")
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Frontend: ").bold = True
    ul.add_run("A React application (using Vite) that provides a user interface to input code, visualize the compilation pipeline, and inspect the outputs of each phase (Tokens, AST, Symbol Table, TAC, and Execution).")

    # Section 3: Phases
    doc.add_heading('3. Compiler Phases', level=1)
    phases = [
        ("Lexical Analysis", "Converts the raw source code into a stream of tokens, discarding whitespace and comments."),
        ("Syntax Analysis", "Parses the tokens into an Abstract Syntax Tree (AST) using a recursive descent algorithm."),
        ("Semantic Analysis", "Checks for variable declarations and scope rules, populating the Symbol Table and assigning unique identifiers for variable shadowing."),
        ("Intermediate Code Generation", "Traverses the AST to generate Three-Address Code (TAC), an intermediate representation close to assembly."),
        ("Optimization", "Applies five iterative passes to the TAC: Constant Folding, Constant Propagation, Algebraic Simplification, Dead Code Elimination, and Temp Forwarding."),
        ("Execution", "A Virtual Machine interprets the optimized TAC to produce final console output, supporting control flow instructions like ifFalse and goto.")
    ]
    for name, desc in phases:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{name}: ").bold = True
        p.add_run(desc)

    doc.add_page_break()

    # Section 4: TAC & Optimization Examples
    doc.add_heading('4. TAC & Optimization Examples', level=1)
    doc.add_paragraph("Example 1: Basic Arithmetic")
    doc.add_paragraph("Source:")
    doc.add_paragraph("int main() {\n    int x = 2 + 3 * 4;\n    print(x);\n    return 0;\n}", style='Intense Quote')
    
    doc.add_paragraph("Unoptimized TAC:")
    doc.add_paragraph("t1 = 3 * 4\nt2 = 2 + t1\nx = t2\nprint x\nreturn 0", style='Intense Quote')

    doc.add_paragraph("Optimized TAC:")
    doc.add_paragraph("print 14\nreturn 0", style='Intense Quote')

    # Section 5: Conclusion
    doc.add_heading('5. Conclusion', level=1)
    doc.add_paragraph("The Mini C Compiler successfully demonstrates the core concepts of compiler design. Through building a complete pipeline from scratch, the intricacies of managing scopes, generating intermediate representations, and performing data-flow based optimizations were thoroughly explored and understood.")

    os.makedirs('docs', exist_ok=True)
    doc.save('docs/Mini_C_Compiler_Report.docx')
    print("Generated Report: docs/Mini_C_Compiler_Report.docx")

def generate_presentation():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Mini C Compiler"
    slide.placeholders[1].text = "Compiler Design (3170701)\nC. K. Pithawala College of Engineering and Technology"

    # Slide 2: Objective
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Objective"
    tf = slide.placeholders[1].text_frame
    tf.text = "To build a complete Mini C Compiler from scratch"
    tf.add_paragraph().text = "Implements all core compilation phases"
    tf.add_paragraph().text = "Visualizes data structures (AST, Symbol Table)"
    tf.add_paragraph().text = "Focuses on Intermediate Code (TAC) and Optimization"

    # Slide 3: Architecture
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Architecture"
    tf = slide.placeholders[1].text_frame
    tf.text = "Frontend: React UI"
    tf.add_paragraph().text = "Interactive code editor"
    tf.add_paragraph().text = "Phase visualization tabs"
    
    p = tf.add_paragraph()
    p.text = "Backend: Python FastAPI"
    p.level = 0
    tf.add_paragraph().text = "Hand-written recursive descent parser"
    tf.add_paragraph().text = "Five-pass TAC Optimizer"
    tf.add_paragraph().text = "TAC Virtual Machine for execution"

    # Slide 4: Optimization Engine
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Optimization Engine"
    tf = slide.placeholders[1].text_frame
    tf.text = "Performs iterative passes over TAC:"
    tf.add_paragraph().text = "1. Constant Folding"
    tf.add_paragraph().text = "2. Constant Propagation"
    tf.add_paragraph().text = "3. Algebraic Simplification"
    tf.add_paragraph().text = "4. Dead Code Elimination"
    tf.add_paragraph().text = "5. Temp Forwarding"

    os.makedirs('docs', exist_ok=True)
    prs.save('docs/Mini_C_Compiler_Presentation.pptx')
    print("Generated Presentation: docs/Mini_C_Compiler_Presentation.pptx")

if __name__ == "__main__":
    generate_report()
    generate_presentation()
